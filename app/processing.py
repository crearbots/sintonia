"""
Lógica de procesamiento de archivos Excel según las reglas de negocio.
Versión robusta + seguimiento de No instalada.
"""
import re
import uuid
from typing import Any, Dict, List
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from datetime import datetime

from .models import Persona, Carga


FUENTES_VALIDAS = [
    "Listado presencial/Punto de información",
    "Censo virtual",
]
FUENTE_PRESENCIAL_ANTERIOR = "Listado de reunión presencial"


def nombre_fuente(valor: str | None) -> str:
    if valor == FUENTE_PRESENCIAL_ANTERIOR:
        return FUENTES_VALIDAS[0]
    return valor or ""

EQUIPOS_CARGA = [
    "Coordinación",
    "Gestión Documental",
    "Formación y Capacitación",
]

MOTIVOS_NO_INSTALADA = [
    "Código de verificación no llega",
    "No trajo el celular",
    "Almacenamiento lleno",
    "Menor de edad (app solo para mayores)",
    "Otro",
]


def normalizar_celular(valor: Any) -> str | None:
    if valor is None:
        return None
    texto = str(valor).strip()
    if not texto or texto.lower() in ("nan", "none", ""):
        return None
    digitos = re.sub(r"\D", "", texto)
    if len(digitos) == 12 and digitos.startswith("57"):
        digitos = digitos[2:]
    if len(digitos) == 10:
        return digitos
    return None


def normalizar_nombre(valor: Any) -> str:
    if valor is None:
        return ""
    texto = re.sub(r"\s+", " ", str(valor).strip())
    if not texto:
        return ""
    return " ".join(p.capitalize() for p in texto.split(" "))


def normalizar_estado(valor: Any) -> str | None:
    """Acepta Instalada o No instalada."""
    if valor is None:
        return None
    texto = str(valor).strip().lower()
    if texto in ("instalada", "instalado"):
        return "Instalada"
    if texto in ("no instalada", "no instalado", "noinstalada"):
        return "No instalada"
    return None


def mapear_columnas(headers: List) -> Dict[str, int]:
    mapping = {}
    for idx, h in enumerate(headers):
        if h is None:
            continue
        h_lower = str(h).strip().lower()
        if h_lower in ("nombre", "hnos/hnas", "hnos", "hnas", "name"):
            mapping["nombre"] = idx
        elif h_lower in ("celular", "telefono", "teléfono", "tel", "phone", "móvil", "movil"):
            mapping["celular"] = idx
        elif h_lower in ("estado", "status", "tiene app", "instalada"):
            mapping["estado"] = idx
    return mapping


def _placeholder_celular() -> str:
    return f"T{uuid.uuid4().hex[:14]}"


def procesar_excel(
    file_path: str,
    fuente: str,
    nombre_archivo: str,
    fecha_listado,
    db: Session,
    cargado_por: str | None = None,
) -> Dict[str, Any]:
    resumen = {
        "nuevos": 0,
        "actualizados": 0,
        "pendientes": 0,
        "seguimiento_no_instalada": 0,
        "ignorados": 0,
        "duplicados_en_archivo": 0,
        "no_revertidos": 0,  # ya eran Instalada y llegó No instalada
        "errores": [],
        "total_filas": 0,
    }

    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
    except Exception as e:
        resumen["errores"].append(f"No se pudo leer el archivo Excel: {str(e)}")
        return resumen

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        resumen["errores"].append("El archivo está vacío.")
        return resumen

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    col_map = mapear_columnas(headers)

    if "nombre" not in col_map or "celular" not in col_map or "estado" not in col_map:
        resumen["errores"].append(
            "No se encontraron las columnas obligatorias: Nombre, Celular y Estado. "
            f"Columnas detectadas: {headers}"
        )
        return resumen

    idx_nombre = col_map["nombre"]
    idx_celular = col_map["celular"]
    idx_estado = col_map["estado"]
    data_rows = rows[1:]
    resumen["total_filas"] = len(data_rows)

    celulares_en_archivo: set[str] = set()

    for row_num, row in enumerate(data_rows, start=2):
        try:
            raw_nombre = row[idx_nombre] if idx_nombre < len(row) else None
            raw_celular = row[idx_celular] if idx_celular < len(row) else None
            raw_estado = row[idx_estado] if idx_estado < len(row) else None

            nombre = normalizar_nombre(raw_nombre)
            celular = normalizar_celular(raw_celular)
            estado = normalizar_estado(raw_estado)

            if estado is None:
                resumen["ignorados"] += 1
                continue

            if not nombre:
                resumen["ignorados"] += 1
                continue

            # Duplicado en el mismo archivo (solo si hay celular válido)
            if celular and celular in celulares_en_archivo:
                resumen["duplicados_en_archivo"] += 1
                continue
            if celular:
                celulares_en_archivo.add(celular)

            existente = None
            if celular:
                existente = db.query(Persona).filter(Persona.celular == celular).first()

            # ---------- INSTALADA ----------
            if estado == "Instalada":
                if celular is None:
                    persona = Persona(
                        celular=_placeholder_celular(),
                        nombre=nombre,
                        estado="Instalada",
                        fuente_ultima=fuente,
                        fecha_listado=fecha_listado,
                        pendiente_revision=True,
                        motivo=None,
                        motivo_detalle=None,
                    )
                    db.add(persona)
                    resumen["pendientes"] += 1
                    continue

                if existente:
                    existente.nombre = nombre
                    existente.fuente_ultima = fuente
                    # No pisar fecha_listado si ya existe (trazabilidad de la gráfica).
                    # Si pasa de No instalada / pendiente a Instalada y no tenía fecha, sí se asigna.
                    if not existente.fecha_listado:
                        existente.fecha_listado = fecha_listado
                    existente.fecha_ultima_carga = datetime.utcnow()
                    existente.estado = "Instalada"
                    existente.pendiente_revision = False
                    existente.motivo = None
                    existente.motivo_detalle = None
                    resumen["actualizados"] += 1
                else:
                    db.add(Persona(
                        celular=celular,
                        nombre=nombre,
                        estado="Instalada",
                        fuente_ultima=fuente,
                        fecha_listado=fecha_listado,
                        pendiente_revision=False,
                    ))
                    resumen["nuevos"] += 1
                continue

            # ---------- NO INSTALADA ----------
            if estado == "No instalada":
                # Regla de oro: si ya es Instalada, no revertir
                if existente and existente.estado == "Instalada" and not existente.pendiente_revision:
                    resumen["no_revertidos"] += 1
                    continue

                if existente:
                    # Actualizar seguimiento existente
                    existente.nombre = nombre
                    existente.fuente_ultima = fuente
                    if not existente.fecha_listado:
                        existente.fecha_listado = fecha_listado
                    existente.fecha_ultima_carga = datetime.utcnow()
                    existente.estado = "No instalada"
                    # No borramos motivo si ya lo tenía
                    if celular is None:
                        existente.pendiente_revision = True
                    else:
                        existente.pendiente_revision = False
                    resumen["actualizados"] += 1
                    if not existente.motivo:
                        resumen["seguimiento_no_instalada"] += 1
                else:
                    cel = celular if celular else _placeholder_celular()
                    db.add(Persona(
                        celular=cel,
                        nombre=nombre,
                        estado="No instalada",
                        fuente_ultima=fuente,
                        fecha_listado=fecha_listado,
                        pendiente_revision=(celular is None),
                        motivo=None,
                        motivo_detalle=None,
                    ))
                    resumen["seguimiento_no_instalada"] += 1
                continue

        except Exception as e:
            resumen["errores"].append(f"Fila {row_num}: {str(e)}")

    try:
        carga = Carga(
            nombre_archivo=nombre_archivo,
            fuente=fuente,
            fecha_listado=fecha_listado,
            total_registros=resumen["total_filas"],
            nuevos=resumen["nuevos"],
            actualizados=resumen["actualizados"],
            pendientes=resumen["pendientes"],
            cargado_por=cargado_por,
        )
        db.add(carga)
        db.commit()
    except IntegrityError as e:
        db.rollback()
        resumen["errores"].append(
            "Error al guardar (posible duplicado). Ningún cambio de esta carga se aplicó. "
            + str(getattr(e, "orig", e))
        )
        resumen["nuevos"] = 0
        resumen["actualizados"] = 0
        resumen["pendientes"] = 0
        resumen["seguimiento_no_instalada"] = 0
    except Exception as e:
        db.rollback()
        resumen["errores"].append(f"Error al guardar la carga: {str(e)}")
        resumen["nuevos"] = 0
        resumen["actualizados"] = 0
        resumen["pendientes"] = 0
        resumen["seguimiento_no_instalada"] = 0

    return resumen
